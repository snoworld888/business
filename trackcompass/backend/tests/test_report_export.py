"""报告导出契约测试：PPTX / DOCX 必须**真的产出文件**，且不得因字段缺失而 500。

## 背景（2026-09-17）

`export()` 的 PPT 分支写死 `content['meta']['template_name']`，而该键实际在
`content` 的**顶层**（见 `build_report` 返回结构）⇒ 必然 `KeyError`。

该缺陷被「`python-pptx` 未安装 → 直接降级 markdown」**长期掩盖**：
装上依赖后首次真跑才暴露（`{"code":500,"message":"服务器错误: 'template_name'"}`）。
所以本文件的存在意义是：**只要导出还在，就必须有真跑产物的测试**，
否则「依赖没装」会继续把代码缺陷藏起来。

> 同类风险：`python-docx` 必须装 —— 代码走 `from docx import Document`，
> 而 `requirements.txt` 原先注释写的是 `docxtpl`（另一个包），
> 「按注释装完仍不生效」。
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from app.services.report import export

STAMP = "2026-09-17T10:00:00"


def _content(**override) -> dict:
    content = {
        "title": "测试报告",
        "template": "investment",
        "template_name": "投资备忘录",
        "sections": [
            {"title": "执行摘要", "markdown": "正文一", "citations": ["E1"]},
            {"title": "风险与应对", "markdown": "正文二", "citations": []},
        ],
        "meta": {"template_tone": "正式", "evidence_count": 3, "page_count": 1},
    }
    content.update(override)
    return content


# ---------------------------------------------------------------- PPTX

def test_pptx_export_produces_real_file(tmp_path: Path) -> None:
    res = export(_content(), "pptx", STAMP, tmp_path)
    assert res["format"] == "pptx", res
    assert "error" not in res
    path = Path(res["path"])
    assert path.exists() and path.stat().st_size > 0

    from pptx import Presentation

    prs = Presentation(str(path))
    # 封面 1 页 + 每个章节 1 页
    assert len(prs.slides) == len(_content()["sections"]) + 1
    assert "投资备忘录" in prs.slides[0].placeholders[1].text


def test_pptx_export_without_meta_template_name(tmp_path: Path) -> None:
    """回归用例：`template_name` 只在顶层、`meta` 里没有 —— 曾因此 500 KeyError。"""
    content = _content()
    content["meta"] = {"evidence_count": 2}  # 故意去掉 template_tone
    res = export(content, "pptx", STAMP, tmp_path)
    assert res["format"] == "pptx", res
    assert Path(res["path"]).exists()


def test_pptx_export_tolerates_missing_optional_fields(tmp_path: Path) -> None:
    """章节缺 title / markdown、顶层缺 template_name，都不得崩。"""
    res = export(
        {"title": "最小报告", "sections": [{}, {"title": "有标题无正文"}]},
        "pptx", STAMP, tmp_path,
    )
    assert res["format"] == "pptx", res
    assert Path(res["path"]).exists()


# ---------------------------------------------------------------- DOCX

def test_docx_export_produces_real_file(tmp_path: Path) -> None:
    res = export(_content(), "docx", STAMP, tmp_path)
    assert res["format"] == "docx", res
    path = Path(res["path"])
    assert path.exists() and path.stat().st_size > 0

    from docx import Document

    doc = Document(str(path))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "执行摘要" in heads and "风险与应对" in heads
    # 证据引用需落进文档
    assert any("E1" in p.text for p in doc.paragraphs)


def test_docx_export_tolerates_missing_fields(tmp_path: Path) -> None:
    res = export({"title": "T", "sections": [{}]}, "docx", STAMP, tmp_path)
    assert res["format"] == "docx", res
    assert Path(res["path"]).exists()


# ---------------------------------------------------------------- 其他格式

def test_markdown_and_json_produce_files(tmp_path: Path) -> None:
    for fmt, suffix in (("markdown", ".md"), ("json", ".json")):
        res = export(_content(), fmt, STAMP, tmp_path)
        path = Path(res["path"])
        assert path.suffix == suffix and path.exists(), res


def test_html_produces_file(tmp_path: Path) -> None:
    res = export(_content(), "html", STAMP, tmp_path)
    path = Path(res["path"])
    assert path.suffix == ".html" and path.exists()
    assert "执行摘要" in path.read_text(encoding="utf-8")


@pytest.mark.parametrize("fmt", ["pdf", "notion", "doc"])
def test_unsupported_formats_report_error_instead_of_crashing(
    tmp_path: Path, fmt: str,
) -> None:
    """PDF / Notion 目前**不支持** —— 必须返回 error 而不是抛异常或静默产出错文件。

    见 `docs/功能状态总表.md` G-01：设计文档 §5.3 要求「可导出 PPT/Word/Notion」。
    PPT/Word 已可用，PDF/Notion 仍未实现；本用例把现状**钉在测试里**，
    将来实现后此用例会失败，提醒同步更新总表。
    """
    res = export(_content(), fmt, STAMP, tmp_path)
    assert res.get("error"), f"{fmt} 应返回 unsupported 错误，实际 {res}"
    assert res.get("path") is None
